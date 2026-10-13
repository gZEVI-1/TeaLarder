#!/usr/bin/env python3
"""
Преобразование чайной Excel-таблицы в JSON для модели Django Tea.

Что делает скрипт:
  * читает XLSX-файл с чайной таблицей;
  * парсит строки в словари, совместимые с моделью Tea;
  * определяет регион произрастания (поле place_of_harvest);
  * переводит цену из руб/кг в руб/г;
  * умеет выводить как обычный JSON, так и fixture для django loaddata.

Примеры запуска:

    python parser_ds.py "Чаай.xlsx" -o teas_ds.json

    python parser_ds.py "Чаай.xlsx" --fixture --app Catalog -o teas.json

    python parser_ds.py "Чаай.xlsx" --fixture --app Catalog -o ../Catalog/fixtures/teas.json
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from datetime import datetime, timezone

# ---------------------------------------------------------------------------
# Справочники (должны совпадать с models.py)
# ---------------------------------------------------------------------------

TEA_TYPES = {
    "white": "Белый",
    "green": "Зеленый",
    "matcha": "Матча",
    "yellow": "Желтый",
    "oolongNF": "Улун северо-фуцзяньский",
    "oolongSF": "Улун южно-фуцзяньский",
    "oolongGD": "Улун гуаньдунский",
    "oolongTW": "Улун тайваньский",
    "oolongGaba": "Улун габа",
    "red": "Красный",
    "black": "Хэй Ча",
    "puerSHU": "Шу Пуэр",
    "puerSHEN": "Шен Пуэр",
    "herbal": "Травяной",
    "flavored": "Ароматизированный",
    "other": "Другое",
}

DEFAULT_TEA_TYPE = "other"

# Регион по умолчанию, если не удалось определить точнее.
DEFAULT_PLACE_OF_HARVEST = "China"

# ---------------------------------------------------------------------------
# Ограничения DecimalField / PositiveIntegerField из models.py
# ---------------------------------------------------------------------------

# price = DecimalField(max_digits=6, decimal_places=2)  -> максимум 9999.99
PRICE_MAX = Decimal("9999.99")

# mark = DecimalField(max_digits=4, decimal_places=1), 0.0..10.0
MARK_MIN = Decimal("0.0")
MARK_MAX = Decimal("10.0")

# year_of_harvest = PositiveIntegerField, 1900..2100
YEAR_MIN = 1900
YEAR_MAX = 2100

# Ожидаемые заголовки колонок в Excel.
EXPECTED_HEADERS = {
    "id": "№ чая",
    "availability": "Наличие",
    "tea_type": "Тип чая",
    "name": "Название",
    "place_of_purchase": "Место покупки",
    "description": "Комментарий",
    "price": "Цена",
    "mark": "Оценка",
}


# ---------------------------------------------------------------------------
# Общие хелперы
# ---------------------------------------------------------------------------

def clean_text(value: Any, *, compact: bool = False) -> str | None:
    """
    Очищает значение ячейки.

    Убирает:
      * неразрывные пробелы;
      * табы;
      * управляющие символы (категория Unicode Cf), например U+200E;
      * лишние пробелы;
      * пустые строки.

    По умолчанию сохраняет переносы строк.
    В compact-режиме переносы превращаются в обычный пробел.
    """
    if value is None:
        return None

    # NBSP -> обычный пробел, таб -> пробел.
    text = str(value).replace("\u00a0", " ").replace("\t", " ")

    # Убираем, например, left-to-right mark (U+200E).
    text = "".join(
        char for char in text
        if unicodedata.category(char) != "Cf"
    ).strip()

    if not text:
        return None

    # Схлопываем подряд идущие пробелы в каждой строке.
    lines = [
        re.sub(r" +", " ", line).strip()
        for line in text.splitlines()
    ]
    lines = [line for line in lines if line]

    if not lines:
        return None

    separator = " " if compact else "\n"
    return separator.join(lines)


def normalize_header(value: Any) -> str:
    """Нормализует заголовок колонки для сравнения."""
    text = clean_text(value, compact=True) or ""
    return re.sub(r"\s+", " ", text.casefold())


# ---------------------------------------------------------------------------
# Парсинг отдельных полей
# ---------------------------------------------------------------------------

def parse_bool(value: Any) -> bool:
    """
    Преобразует значение колонки 'Наличие' в bool.

    '+' / 'да' / 'true' / 1 -> True.
    Всё остальное            -> False.
    """
    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        return bool(value)

    text = clean_text(value, compact=True)

    if text is None:
        return False

    return text.casefold() in {"+", "да", "true", "1"}


def normalize_tea_type(value: Any) -> tuple[str, str | None]:
    """
    Преобразует текстовый тип чая в код из TEA_TYPES.

    Возвращает пару:
        (код типа, текст предупреждения или None)
    """
    raw = clean_text(value, compact=True)

    if not raw:
        return DEFAULT_TEA_TYPE, "тип чая не указан"

    text = raw.casefold().replace("ё", "е")
    compact = re.sub(r"\s+", "", text)

    # --- Улуны ---
    if "улун" in text and "габа" in text:
        return "oolongGaba", None

    if "улун" in text:
        # Пометки в скобках после слова "Улун": ГД / ТВ / ЮФ / УИ / КТ.
        if "гд" in compact or "гуаньдун" in text or "гуандун" in text:
            return "oolongGD", None

        if "тв" in compact or "тайван" in text:
            return "oolongTW", None

        if "юф" in compact or "южно" in text:
            return "oolongSF", None

        # Если явных пометок нет, считаем северо-фуцзяньским.
        return "oolongNF", None

    # --- Пуэры (порядок важен: "шу пуэр" раньше, чем "шен пуэр") ---
    if "шу" in text and "пуэр" in text:
        return "puerSHU", None

    if ("шен" in text or "шэн" in text) and "пуэр" in text:
        return "puerSHEN", None

    # --- Остальные категории ---
    if "хей" in text or "хэй" in text:
        return "black", None

    if "матча" in text or "маття" in text:
        return "matcha", None

    if "бел" in text:
        return "white", None

    if "зелен" in text:
        return "green", None

    if "желт" in text:
        return "yellow", None

    if "красн" in text:
        return "red", None

    if "трав" in text:
        return "herbal", None

    if "аромат" in text:
        return "flavored", None

    return DEFAULT_TEA_TYPE, f"неизвестный тип чая {raw!r}"


def detect_region(
    tea_type: str,
    raw_type: Any,
    name: Any,
    description: Any = None,
    place_of_purchase: Any = None,
) -> str:
    """
    Определяет регион произрастания — код из REGIONS в models.py.

    Правила:
      * Пуэры (шу и шен) -> Yunnan.
      * Улуны:
          - пометка ГД  -> Guangdong;
          - пометка ТВ  -> Taiwan;
          - пометка ЮФ  -> Fujian;
          - пометка УИ  -> Fujian (Уишань);
          - пометка КТ  -> Taiwan;
          - конкретные названия сортов (Те Гуань Инь, Да Хун Пао и т.д.).
      * Для остальных типов — пытаемся найти регион в тексте.
      * Если ничего не нашли — возвращаем China.
    """
    raw_type_text = clean_text(raw_type, compact=True) or ""
    name_text = clean_text(name, compact=True) or ""
    description_text = clean_text(description, compact=True) or ""
    purchase_text = clean_text(place_of_purchase, compact=True) or ""

    raw_type_lower = raw_type_text.casefold().replace("ё", "е")
    raw_type_compact = re.sub(r"\s+", "", raw_type_lower)

    name_lower = name_text.casefold().replace("ё", "е")
    description_lower = description_text.casefold().replace("ё", "е")
    purchase_lower = purchase_text.casefold().replace("ё", "е")

    # --- Пуэры ---
    # И шу, и шен традиционно делают в Юньнани.
    if tea_type in {"puerSHU", "puerSHEN"}:
        return "Yunnan"

    # --- Улуны ---
    if tea_type.startswith("oolong"):
        # Явные пометки в колонке "Тип чая".
        if "гд" in raw_type_compact or "гуандун" in raw_type_lower:
            return "Guangdong"

        if "тв" in raw_type_compact or "тайван" in raw_type_lower:
            return "Taiwan"

        if "юф" in raw_type_compact or (
            "южно" in raw_type_lower and "фуцзянь" in raw_type_lower
        ):
            return "Fujian"

        if "уи" in raw_type_compact or "уишань" in raw_type_lower:
            return "Fujian"

        if "кт" in raw_type_compact:
            return "Taiwan"

        # Определение по названию конкретного сорта.
        # Фуцзянь (север + юг).
        if re.search(
            r"те\s*гуань\s*инь|тегуаньинь|тигуаньинь|те\s*гуфнь",
            name_lower,
        ):
            return "Fujian"

        if re.search(
            r"да\s*хун\s*пао|дахунпао|шуй\s*сянь|шуйсянь|"
            r"жоу\s*гуй|жоугуй",
            name_lower,
        ):
            return "Fujian"

        # Гуандун (дан цуны).
        if re.search(
            r"фен\s*хуан|фэн\s*хуан|дан\s*цун|данцун",
            name_lower,
        ):
            return "Guangdong"

        # Тайвань.
        if re.search(
            r"дун\s*дин|дундин|алишань|лишань|"
            r"наньтоу|бао\s*чжун|баочжун|габа",
            name_lower,
        ):
            return "Taiwan"

        # Запасной вариант — по коду типа.
        return {
            "oolongNF": "Fujian",
            "oolongSF": "Fujian",
            "oolongGD": "Guangdong",
            "oolongTW": "Taiwan",
            "oolongGaba": "Taiwan",
        }.get(tea_type, DEFAULT_PLACE_OF_HARVEST)

    # --- Остальные типы ---
    # Пытаемся найти регион в тексте.
    combined = f"{name_lower} {description_lower} {purchase_lower}"

    if "юньнань" in combined or "мэнхай" in combined or "махэй" in combined:
        return "Yunnan"
    if "тайвань" in combined:
        return "Taiwan"
    if "фуцзянь" in combined:
        return "Fujian"
    if "гуандун" in combined:
        return "Guangdong"

    return DEFAULT_PLACE_OF_HARVEST


def parse_decimal(value: Any) -> Decimal | None:
    """
    Парсит число.

    Поддерживает:
        8.5
        8,5
        1 234.56
        1 234,56
        ="8,8"
    """
    if value is None or value == "":
        return None

    if isinstance(value, Decimal):
        return value

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None

    text = clean_text(value, compact=True)

    if text is None:
        return None

    # Excel-формула вида ="8,8".
    if text.startswith("="):
        text = text.removeprefix("=").strip().strip('"')

    match = re.fullmatch(
        r"([+-]?\d+(?:[ ]?\d{3})*)(?:[.,](\d+))?",
        text,
    )

    if not match:
        return None

    integer_part = match.group(1).replace(" ", "")
    fraction_part = match.group(2)

    if fraction_part is None:
        return Decimal(integer_part)

    return Decimal(f"{integer_part}.{fraction_part}")


def parse_price(value: Any) -> Decimal | None:
    """
    Парсит цену и возвращает её в руб/кг.

    Примеры:
        4957 руб/кг          -> 4957.00
        13750руб/кг          -> 13750.00
        бесплатно            -> 0.00
        21190 руб/кг (6880 руб/кг)
                             -> 6880.00   (берём последнее число —
                                            обычно это актуальная цена)

    Внимание: перевод в руб/г делается отдельно в parse_teas.
    """
    if value is None or value == "":
        return None

    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = parse_decimal(value)

        if number is None:
            return None

        return number.quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP,
        )

    text = clean_text(value, compact=True)

    if text is None:
        return None

    lower_text = text.casefold()

    if "бесплатно" in lower_text or "даром" in lower_text:
        return Decimal("0.00")

    matches = re.findall(
        r"(?<![\d.,])"
        r"(\d+(?:[ ]\d{3})*(?:[.,]\d+)?)"
        r"(?![\d.,])",
        text,
    )

    if not matches:
        return None

    # Если в тексте несколько чисел — берём последнее
    # (обычно это цена в скобках, то есть актуальная).
    number = parse_decimal(matches[-1])

    if number is None:
        return None

    return number.quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )


def parse_mark(value: Any) -> Decimal | None:
    """
    Парсит оценку и приводит её к одному знаку после запятой.

    Пустое значение и дефис преобразуются в None.
    """
    if isinstance(value, str):
        text = clean_text(value, compact=True)

        if text in {"-", "—"}:
            return None

    number = parse_decimal(value)

    if number is None:
        return None

    return number.quantize(
        Decimal("0.1"),
        rounding=ROUND_HALF_UP,
    )


def extract_year(*values: Any) -> int | None:
    """
    Ищет год сначала в названии, затем в комментарии.

    Например:
        Красный бык 2022       -> 2022
        Булан 2012             -> 2012
        весна 2024             -> 2024
    """
    for value in values:
        text = clean_text(value)

        if not text:
            continue

        match = re.search(
            r"(?<!\d)((?:19|20)\d{2})(?!\d)",
            text,
        )

        if match:
            return int(match.group(1))

    return None


# ---------------------------------------------------------------------------
# Разбор листа Excel
# ---------------------------------------------------------------------------

def build_columns(
    ws: Worksheet,
    header_row: int,
) -> dict[str, int]:
    """
    Находит номера колонок по заголовкам.

    Например:
        {
            "id": 2,
            "availability": 3,
            "tea_type": 4,
            ...
        }
    """
    normalized_to_column: dict[str, int] = {}

    for column in range(1, ws.max_column + 1):
        normalized = normalize_header(
            ws.cell(header_row, column).value
        )

        if normalized:
            normalized_to_column[normalized] = column

    columns: dict[str, int] = {}

    for field, expected_header in EXPECTED_HEADERS.items():
        normalized_header = normalize_header(expected_header)
        column = normalized_to_column.get(normalized_header)

        if column is None:
            raise ValueError(
                f"В строке {header_row} не найдена колонка "
                f"{expected_header!r}"
            )

        columns[field] = column

    return columns


def decimal_to_json(
    value: Decimal | None,
    *,
    as_string: bool,
) -> float | str | None:
    """Преобразует Decimal в JSON-совместимое значение."""
    if value is None:
        return None

    if as_string:
        return format(value, "f")

    return float(value)


def parse_teas(
    ws: Worksheet,
    *,
    header_row: int = 4,
    decimals_as_strings: bool = False,
) -> tuple[list[dict[str, Any]], list[str]]:
    """
    Парсит чаи с листа Excel.

    Возвращает:
        (список чаев, список предупреждений)
    """
    columns = build_columns(ws, header_row)

    teas: list[dict[str, Any]] = []
    issues: list[str] = []

    for row_number in range(header_row + 1, ws.max_row + 1):
        raw_id = ws.cell(
            row_number,
            columns["id"],
        ).value

        raw_name = ws.cell(
            row_number,
            columns["name"],
        ).value

        name = clean_text(raw_name, compact=True)

        # Пропускаем полностью пустые строки.
        if raw_id is None and name is None:
            continue

        # Пропускаем строки-заготовки, где указан только номер.
        if name is None:
            issues.append(
                f"строка {row_number}: пропущена, "
                f"потому что не заполнено название"
            )
            continue

        # --- id ---
        try:
            tea_id = int(str(raw_id).strip())
        except (TypeError, ValueError):
            issues.append(
                f"строка {row_number}: некорректный номер чая "
                f"{raw_id!r}"
            )
            continue

        if not 1 <= tea_id <= 9999:
            issues.append(
                f"строка {row_number}: id={tea_id} "
                f"вне диапазона 1..9999"
            )

        # --- тип чая ---
        raw_type = ws.cell(
            row_number,
            columns["tea_type"],
        ).value

        tea_type, type_issue = normalize_tea_type(raw_type)

        if type_issue:
            issues.append(
                f"строка {row_number}: {type_issue}; "
                f"использован тип 'other'"
            )

        # --- место покупки ---
        raw_place_of_purchase = ws.cell(
            row_number,
            columns["place_of_purchase"],
        ).value

        place_of_purchase = clean_text(
            raw_place_of_purchase,
            compact=True,
        ) or ""

        # --- комментарий ---
        raw_description = ws.cell(
            row_number,
            columns["description"],
        ).value

        description = clean_text(raw_description) or ""

        # --- цена ---
        # В таблице цены указаны в руб/кг.
        # В модели price — DecimalField(max_digits=6, decimal_places=2).
        # Переводим в руб/г: price = цена_за_кг / 1000.
        raw_price = ws.cell(
            row_number,
            columns["price"],
        ).value

        price_per_kg = parse_price(raw_price)

        if price_per_kg is None:
            price = None
        else:
            price = (
                price_per_kg / Decimal("1000")
            ).quantize(
                Decimal("0.01"),
                rounding=ROUND_HALF_UP,
            )

        if price is not None and price > PRICE_MAX:
            issues.append(
                f"строка {row_number}: price={price} руб/г "
                f"не помещается в DecimalField("
                f"max_digits=6, decimal_places=2); "
                f"максимум {PRICE_MAX}"
            )

        # --- оценка ---
        raw_mark = ws.cell(
            row_number,
            columns["mark"],
        ).value

        mark = parse_mark(raw_mark)

        if mark is not None and not MARK_MIN <= mark <= MARK_MAX:
            issues.append(
                f"строка {row_number}: mark={mark} "
                f"вне диапазона 0..10"
            )

        # --- год сбора ---
        year = extract_year(raw_name, raw_description)

        if year is not None and not YEAR_MIN <= year <= YEAR_MAX:
            issues.append(
                f"строка {row_number}: "
                f"year_of_harvest={year} вне диапазона 1900..2100"
            )
            year = None

        # --- регион произрастания ---
        # Пуэры -> Yunnan.
        # Улуны -> по пометкам (ГД/ТВ/ЮФ/УИ/КТ) и по названию сорта.
        # Остальные — по тексту или China.
        place_of_harvest = detect_region(
            tea_type=tea_type,
            raw_type=raw_type,
            name=raw_name,
            description=raw_description,
            place_of_purchase=raw_place_of_purchase,
        )

        # --- проверки длины строк ---
        if len(name) > 200:
            issues.append(
                f"строка {row_number}: название длиннее 200 символов"
            )

        if len(place_of_purchase) > 200:
            issues.append(
                f"строка {row_number}: место покупки длиннее 200 символов"
            )

        # --- наличие ---
        availability = parse_bool(
            ws.cell(
                row_number,
                columns["availability"],
            ).value
        )

        teas.append(
            {
                "id": tea_id,
                "availability": availability,
                "name": name,
                "place_of_purchase": place_of_purchase,
                "description": description,
                "tea_type": tea_type,
                "price": decimal_to_json(
                    price,
                    as_string=decimals_as_strings,
                ),
                "mark": decimal_to_json(
                    mark,
                    as_string=decimals_as_strings,
                ),
                "year_of_harvest": year,
                "place_of_harvest": place_of_harvest,
            }
        )

    return teas, issues


# ---------------------------------------------------------------------------
# Fixture для django loaddata
# ---------------------------------------------------------------------------
def make_fixture(
    teas: list[dict[str, Any]],
    app_label: str,
    created_at: datetime | None = None,
) -> list[dict[str, Any]]:
    """
    Преобразует обычный JSON в формат fixture для django loaddata.

    Добавляет created_at и updated_at, потому что loaddata вызывает
    Model.save_base(raw=True), и auto_now_add / auto_now не срабатывают.
    """
    timestamp = created_at or datetime.now(timezone.utc)
    iso_timestamp = timestamp.astimezone(timezone.utc).isoformat().replace(
        "+00:00", "Z",
    )

    return [
        {
            "model": f"{app_label}.Tea",
            "pk": tea["id"],
            "fields": {
                **{
                    key: value
                    for key, value in tea.items()
                    if key != "id"
                },
                "created_at": iso_timestamp,
                "updated_at": iso_timestamp,
            },
        }
        for tea in teas
    ]
# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Преобразование чайной Excel-таблицы в JSON "
            "для модели Django Tea."
        )
    )

    


    parser.add_argument(
        "input",
        type=Path,
        help="путь к XLSX-файлу",
    )

    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="путь к JSON-файлу; без флага JSON выводится в консоль",
    )

    parser.add_argument(
        "--sheet",
        help="имя листа; по умолчанию используется первый лист",
    )

    parser.add_argument(
        "--header-row",
        type=int,
        default=4,
        help="строка с заголовками чаев, по умолчанию 4",
    )

    parser.add_argument(
        "--fixture",
        action="store_true",
        help="сделать JSON в формате django loaddata",
    )

    parser.add_argument(
        "--app",
        help="label приложения Django для --fixture, например: tea",
    )

    parser.add_argument(
        "--decimals-as-strings",
        action="store_true",
        help=(
            "сохранять DecimalField как строки, "
            "например '4957.00'"
        ),
    )

    parser.add_argument(
        "--strict",
        action="store_true",
        help="завершаться с ошибкой при наличии предупреждений",
    )

    args = parser.parse_args()

    if args.fixture and not args.app:
        parser.error(
            "для --fixture обязательно укажите --app"
        )

    if not args.input.is_file():
        parser.error(
            f"файл не найден: {args.input}"
        )

    workbook = load_workbook(
        args.input,
        read_only=False,
        data_only=True,
    )

    if args.sheet:
        worksheet = workbook[args.sheet]
    else:
        worksheet = workbook.worksheets[0]

    teas, issues = parse_teas(
        worksheet,
        header_row=args.header_row,
        decimals_as_strings=args.decimals_as_strings,
    )

    if args.fixture:
        data: Any = make_fixture(
            teas,
            args.app,
        )
    else:
        data = teas

    json_text = json.dumps(
        data,
        ensure_ascii=False,
        indent=2,
    )

    if args.output:
        args.output.write_text(
            json_text + "\n",
            encoding="utf-8",
        )
    else:
        print(json_text)

    for issue in issues:
        print(
            f"warning: {issue}",
            file=sys.stderr,
        )

    if issues and args.strict:
        print(
            f"error: обнаружено предупреждений: {len(issues)}",
            file=sys.stderr,
        )
        return 1

    print(
        f"Обработано чаев: {len(teas)}",
        file=sys.stderr,
    )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())